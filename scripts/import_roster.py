import openpyxl,json,hashlib,pathlib,sys
source=sys.argv[1]
w=openpyxl.load_workbook(source,read_only=True,data_only=True)
# Та же формула, что rosterId/studentId в src/domain.js – менять только вместе.
def uid(prefix,value): return prefix+hashlib.sha256(value.encode()).hexdigest()[:16]
students={}; teachers={}; enrollments=[]
rows=[r for r in list(w['База'].values)[1:] if all(r[:4])]
# «Вак_Иванов» и «Иванов» – один человек; при наличии обоих написаний берём без приставки.
names={str(r[1]).strip() for r in rows}
merged={n:n[4:] for n in names if n.startswith('Вак_') and n[4:] in names}
for row in rows:
 group,teacher,student,course,kind,*_=row
 teacher=merged.get(str(teacher).strip(),teacher)
 sid=uid('s_',str(student).strip());tid=uid('t_',str(teacher).strip())
 students[sid]={'id':sid,'name':str(student).strip()}
 teachers[tid]={'id':tid,'name':str(teacher).strip()}
 for g in str(group).split(';'):
  enrollments.append({'studentId':sid,'teacherId':tid,'group':g.strip(),'course':course,'kind':kind})
result={'students':list(students.values()),'teachers':list(teachers.values()),'enrollments':enrollments,'quality':{'sourceRows':w['База'].max_row-1,'unresolved':w['Не определены'].max_row-1,'mismatches':w['Несовпадения'].max_row-1,'importedAt':__import__('datetime').datetime.now(__import__('datetime').timezone.utc).isoformat()}}
pathlib.Path('data').mkdir(exist_ok=True)
pathlib.Path('data/roster.json').write_text(json.dumps(result,ensure_ascii=False),encoding='utf-8')
print(json.dumps({k:len(result[k]) for k in ['students','teachers','enrollments']}))
