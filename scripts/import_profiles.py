"""Заносит в базу журнала сведения о студентах из списков международного офиса.

Источники: список иностранных студентов факультета (программа, курс, гражданство)
и таблицы миграционного учёта (куратор, проживание, основание пребывания, сроки).
Заполняются только пустые поля карточек и отсутствующие процедуры: правки,
внесённые через сайт, не затираются. Без --apply ничего не пишет, только отчёт.

uv run --with openpyxl python scripts/import_profiles.py data/attendance.sqlite \
    --list "Список ИС ФП.xlsx" --migration Книга1.xlsx --migration "Не требуют контроля.xlsx" --apply
Сервер на время записи остановить: новых студентов он увидит после запуска.
"""
import argparse, datetime, difflib, hashlib, json, re, sqlite3
import openpyxl

PROGRAMS = {
    "Право", "Публичное право", "Юрист в бизнесе", "Юриспруденция",
    "Юриспруденция: цифровой юрист", "Цифровой юрист",
    "Комплаенс и профилактика правовых рисков",
    "Комплаенс и профилактика правовых рисков (совместная с ЧГУ)",
    "Сравнительное правоведение и фундаментальный правовой анализ",
    "Цифровое право", "Современное частное право", "Фармправо и здравоохранение",
    "Юрист в правосудии", "Право международной торговли и разрешение споров", "ЛигалТех",
}  # src/office.js
RESIDENCE = {"виза": "visa", "безвиз": "visa_free", "безвизРБ": "visa_free",
             "ВНЖ": "residence_permit", "РВП": "rvp", "РВП истёкбезвиз": "visa_free"}
HOUSING = {"Общежитие": "dormitory", "Частный адрес": "private"}


def clean(name):
    return re.sub(r"\s+", " ", re.sub(r"\s-$", "", str(name).strip()))


def norm(name):
    return clean(name).lower().replace("ё", "е")


def date(cell):
    m = re.match(r"(\d{2})\.(\d{2})\.(\d{4})", str(cell or ""))
    return f"{m[3]}-{m[2]}-{m[1]}" if m else ""


def read_list(path):
    rows = list(openpyxl.load_workbook(path, data_only=True).active.iter_rows(values_only=True))[1:]
    out = {}
    for fio, latin, version, course, citizenship, sending in rows:
        if not fio:
            continue
        version = clean(version or "")
        program = version.split(" очная ", 1)[-1] if " очная " in version else ""
        out[norm(fio)] = {
            "name": clean(fio),
            "nameLatin": clean(latin or ""),
            "programVersion": version,
            "program": program if program in PROGRAMS else "",
            "year": int(re.sub(r"\D", "", course or "") or 0),
            "citizenship": clean(citizenship or ""),
            "sendingCountry": clean(sending or ""),
        }
    return out


def read_migration(path):
    out = {}
    for row in openpyxl.load_workbook(path, data_only=True).active.iter_rows(values_only=True):
        row = list(row[1:] if row[0] is None else row[:13]) + [None] * 13
        fio, curator, housing, citizenship, in_russia, status = row[:6]
        if not (fio and curator and citizenship):
            continue  # служебная строка под студентом
        out[norm(fio)] = {
            "name": clean(fio), "curator": clean(curator), "housing": HOUSING.get(housing, ""),
            "citizenship": clean(citizenship).title(),
            "inRussia": {"да": "yes", "нет": "no"}.get(in_russia, ""),
            "status": status or "", "russian": str(status).startswith("гражд. РФ"),
            "registration": str(row[7] or ""), "visa": str(row[8] or ""),
            "migrationCardUntil": date(row[9]), "passportUntil": date(row[10]),
            "source": path.rsplit("/", 1)[-1],
        }
    return out


def procedures(m, today):
    stamp = f"Источник: таблица миграционного учёта «{m['source']}», внесено {today}."
    if m["inRussia"] == "no":
        note = "Студент не в РФ – контроль не требуется. " + stamp
        return {k: {"state": "exempt", "note": note} for k in ("registration", "visa")}
    out = {}
    if date(m["registration"]):
        extra = re.search(r"\+\d+ дн РБ", m["registration"])
        out["registration"] = {"state": "pending", "dueDate": date(m["registration"]),
                               "note": "Срок миграционного учёта" + (f" ({extra[0]})" if extra else "") + ". " + stamp}
    if m["visa"].startswith("не требуется"):
        out["visa"] = {"state": "exempt", "note": f"Виза не требуется: {m['status']}. " + stamp}
    elif date(m["visa"]):
        out["visa"] = {"state": "pending", "dueDate": date(m["visa"]), "note": "Срок действия визы. " + stamp}
    elif "виза требуется" in m["visa"]:
        out["visa"] = {"state": "pending", "note": "Виза требуется, срока в таблице нет. " + stamp}
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("db")
    ap.add_argument("--list", required=True)
    ap.add_argument("--migration", action="append", default=[])
    ap.add_argument("--apply", action="store_true")
    a = ap.parse_args()
    today = datetime.date.today().isoformat()
    listed, migration = read_list(a.list), {}
    for path in a.migration:
        migration.update(read_migration(path))
    russians = {k for k, m in migration.items() if m["russian"]}

    db = sqlite3.connect(a.db)
    registry = {norm(n): i for i, n in db.execute("SELECT id,name FROM roster_students")}

    def match(key):
        """Точное ФИО, затем единственное похожее написание (отчество, буква)."""
        if key in registry:
            return key, False
        close = [k for k in registry if k.split()[:2] == key.split()[:2]
                 or difflib.SequenceMatcher(None, k, key).ratio() >= 0.9]
        return (close[0], True) if len(close) == 1 else (None, False)

    people, fuzzy, added, skipped = {}, [], [], []
    for source in (listed, migration):
        for key, data in source.items():
            if key in russians:
                skipped.append(data["name"])
                continue
            found, approx = match(key)
            if approx:
                fuzzy.append((data["name"], found))
            if not found:
                found = key
                # Та же формула, что studentId в src/domain.js – менять только вместе.
                registry[key] = "s_" + hashlib.sha256(data["name"].encode()).hexdigest()[:16]
                added.append(data["name"])
                db.execute("INSERT OR IGNORE INTO roster_students VALUES(?,?)", (registry[key], data["name"]))
            people.setdefault(found, {}).update({"list" if source is listed else "migration": data})

    filled = created = 0
    for key, src in people.items():
        sid = registry[key]
        row = db.execute("SELECT data FROM student_profiles WHERE studentId=?", (sid,)).fetchone()
        profile = json.loads(row[0]) if row else {}
        l, m = src.get("list", {}), src.get("migration", {})
        wanted = {
            "program": l.get("program", ""), "year": l.get("year", 0),
            "citizenship": l.get("citizenship") or m.get("citizenship", ""),
            "nameLatin": l.get("nameLatin", ""), "sendingCountry": l.get("sendingCountry", ""),
            "programVersion": l.get("programVersion", ""), "curator": m.get("curator", ""),
            "housing": m.get("housing", ""), "inRussia": m.get("inRussia", ""),
            "residence": RESIDENCE.get(m.get("status"), ""),
            "passportUntil": m.get("passportUntil", ""), "migrationCardUntil": m.get("migrationCardUntil", ""),
        }
        before = dict(profile)
        for field, value in wanted.items():
            if value and not profile.get(field):
                profile[field] = value
        if profile.get("foreignStatus", "unknown") == "unknown":
            profile["foreignStatus"] = "confirmed"
        for field, default in (("arrivalDate", ""), ("residence", ""), ("enrollmentStatus", "active"),
                               ("program", ""), ("year", 0), ("citizenship", "")):
            profile.setdefault(field, default)
        if profile != before:
            profile["version"] = before.get("version", 0) + 1
            filled += 1
            db.execute("INSERT OR REPLACE INTO student_profiles VALUES(?,?)",
                       (sid, json.dumps(profile, ensure_ascii=False)))
        for kind, record in (procedures(m, today) if m else {}).items():
            if db.execute("SELECT 1 FROM procedures WHERE studentId=? AND kind=?", (sid, kind)).fetchone():
                continue
            record = {"version": 1, "dueDate": "", "validUntil": "", "completedAt": "", **record,
                      "checkedBy": "Импорт списков", "updatedAt": datetime.datetime.now(datetime.UTC).isoformat()}
            db.execute("INSERT INTO procedures VALUES(?,?,?)", (sid, kind, json.dumps(record, ensure_ascii=False)))
            created += 1

    print(f"В списке факультета: {len(listed)}, в таблицах миграционного учёта: {len(migration)}")
    print(f"Карточек заполнено или дополнено: {filled}; процедур создано: {created}")
    print(f"Новых студентов в реестре (без связей с преподавателями): {len(added)}")
    for n in added:
        print("   +", n)
    print(f"Сопоставлены по похожему написанию ({len(fuzzy)}) – проверьте:")
    for theirs, ours in fuzzy:
        print(f"   {theirs}  →  {ours}")
    print(f"Пропущены граждане РФ ({len(set(skipped))}): {', '.join(sorted(set(skipped)))}")
    no_program = [s["list"]["name"] for s in people.values() if "list" in s and not s["list"]["program"]]
    print(f"Программа не из справочника журнала ({len(no_program)}): {', '.join(no_program)}")
    in_registry_only = [k for k in registry if k not in people]
    print(f"Есть в реестре, но нет ни в одном списке ({len(in_registry_only)}): {', '.join(in_registry_only)}")
    if a.apply:
        db.execute("INSERT INTO audit(actor,action,entity,at) VALUES(?,?,?,?)",
                   ("Импорт списков", "profiles.import", f"{filled} карточек, {created} процедур, {len(added)} новых",
                    datetime.datetime.now(datetime.UTC).isoformat()))
        db.commit()
        print("Записано в базу.")
    else:
        db.rollback()
        print("Сухой прогон: в базу ничего не записано. Для записи добавьте --apply.")


if __name__ == "__main__":
    main()
